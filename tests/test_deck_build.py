"""deck_build: validation messages, film-scene compatibility, reproducible PPTX, lint. Skips when python-pptx is not installed."""
import json, sys
from pathlib import Path
import pytest
pytest.importorskip('pptx'); pytest.importorskip('defusedxml')
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools/deck'))
import deck_build as db

XML = '''<deck title="T" target="google-slides">
<slide layout="title"><eyebrow>Hi</eyebrow><headline>Meet *Studio*</headline><sub>Films and decks</sub><notes>n</notes></slide>
<slide layout="bullets"><heading>Steps</heading><item>One</item><item>Two</item><notes>n</notes></slide>
<slide layout="quote"><quote>The best interface is the one the agent *can build*.</quote><name>Ada</name><role>First programmer</role><notes>n</notes></slide>
<slide layout="stats"><kicker>Numbers</kicker><stat value="9" label="a"/><stat value="98%" label="b"/><notes>n</notes></slide>
</deck>'''

def errs_of(src, **kw): return db.run(src, **kw)['errors']

def test_valid_deck_builds_clean_and_reproducibly(tmp_path):
    a, b = tmp_path / 'a.pptx', tmp_path / 'b.pptx'
    ra, rb = db.run(XML, str(a)), db.run(XML, str(b))
    assert ra['ok'] and not ra['errors'] and ra['lint'] == [] and ra['slides'] == 4
    assert a.read_bytes() == b.read_bytes() and ra['sha256'] == rb['sha256']

def test_exact_messages():
    e = errs_of('<deck><slide layout="title"><headline></headline></slide><slide layout="nope"/></deck>')
    assert 'slide 1 (title): headline is required' in [x['message'] for x in e]
    assert any(x['message'].startswith("slide 2: unknown layout 'nope'; choose one of") for x in e)
    e = errs_of('<deck><slide layout="bullets"><heading>h</heading>' + '<item>x</item>' * 7 + '</slide></deck>')
    assert e[0]['message'] == 'slide 1 (bullets): items has 7 entries, the limit is 6'

def test_film_json_scenes_and_film_only():
    films = json.dumps({'scenes': [{'type': 'title', 'headline': 'Hi'}, {'type': 'steps', 'heading': 'H', 'steps': 'a\nb'},
                                   {'type': 'stats', 'stats': '9 | a\n10 | b'}, {'type': 'chart'}, {'type': 'quote', 'quote': 'q', 'name': 'n'}]})
    e = errs_of(films); assert 'film-only' in e[0]['message']
    r = db.run(films, skip_film_only=True); assert not r['errors'] and r['slides'] == 4
    assert any(w['code'] == 'film-only' for w in r['warnings'])

def test_overlong_value_is_an_error_not_a_bad_slide():
    e = errs_of('<deck><slide layout="stats"><stat value="12345678901" label="x"/></slide></deck>')
    assert 'limit is 10' in e[0]['message']

def test_unknown_target_and_bad_xml():
    assert 'unknown target' in errs_of('<deck target="keynote"><slide layout="title"><headline>x</headline></slide></deck>')[0]['message']
    assert errs_of('<deck><slide>')[0]['code'] == 'xml'

def test_xml_entities_are_refused():
    assert errs_of('<!DOCTYPE d [<!ENTITY x "y">]><deck><slide layout="title"><headline>&x;</headline></slide></deck>')[0]['code'] == 'xml'

def test_schema_lists_layouts():
    s = db.schema(); assert set(s['layouts']) == {'title', 'bullets', 'stats', 'cards', 'table', 'media', 'quote', 'cta'} and 'chart' in s['film_only']

def test_quote_too_long_and_missing():
    e = errs_of('<deck><slide layout="quote"><quote>' + 'x' * 241 + '</quote></slide></deck>'); assert 'limit is 240' in e[0]['message']
    assert 'quote is required' in errs_of('<deck><slide layout="quote"><name>n</name></slide></deck>')[0]['message']

def test_quote_payload_twin_and_live_link(tmp_path):
    x = '<deck target="google-slides"><slide layout="quote"><quote>Hi *there*</quote><name>Ada</name><role>R</role><notes>n</notes></slide></deck>'
    r = db.run(x)
    assert r['payloads'][1] == {'blocks': [{'type': 'quote', 'text': 'Hi there', 'attribution': 'Ada, R'}]}
    out = tmp_path / 'q.pptx'; r = db.run(x, str(out), link_payloads=True)
    assert r['ok'] and r['lint'] == [] and r['live_urls'][1].startswith('https://')
    from pptx import Presentation
    sh = [s for s in Presentation(str(out)).slides[0].shapes if s.name == 'Live link'][0]
    assert sh.click_action.hyperlink.address == r['live_urls'][1]
    assert db.run(x, str(tmp_path / 'p.pptx'))['payloads'] and 'live_urls' not in db.run(x)

def test_media_layout_animated_gif(tmp_path):
    from PIL import Image
    fr = [Image.new('RGB', (160, 90), (i * 40, 30, 60)) for i in range(4)]
    fr[0].save(tmp_path / 'a.gif', save_all=True, append_images=fr[1:], duration=100, loop=0)
    x = tmp_path / 'd.xml'
    x.write_text('<deck target="google-slides"><slide layout="media"><heading>Loop</heading><media src="a.gif"/><alt>Four coloured frames</alt><notes>n</notes></slide></deck>')
    r = db.run(str(x), str(tmp_path / 'o.pptx'))                                   # media paths resolve against the deck file's folder
    assert r['ok'] and r['lint'] == [] and r['media'][1]['frames'] == 4
    noalt = tmp_path / 'n.xml'; noalt.write_text(x.read_text().replace('<alt>Four coloured frames</alt>', ''))
    assert 'alt is required' in db.run(str(noalt))['errors'][0]['message']
    e = db.run('<deck><slide layout="media"><heading>h</heading><media src="nope.gif"/><alt>a</alt></slide></deck>')['errors']
    assert 'file not found: nope.gif' in e[0]['message']

def test_studio_state_becomes_a_deck():
    st = {'name': 'Data story', 'st': {'title': 'Data story', 'preset': 'ocean', 'scenes': [
        {'type': 'title', 'transition': 'dissolve', 'dur': 4.5, 'f': {'eyebrow': 'Q4 REVIEW', 'headline': 'A year\nin *numbers.*', 'sub': ''}},
        {'type': 'chart', 'transition': 'push', 'dur': 5, 'f': {'kind': 'bars'}},
        {'type': 'stats', 'transition': 'slice', 'dur': 5, 'f': {'roll': 'roll', 'kicker': 'BY THE NUMBERS', 'stats': '£2.4M | revenue\n31% | growth\n4,800 | customers'}},
        {'type': 'steps', 'transition': 'wipe', 'dur': 5, 'f': {'heading': 'What\n*comes next.*', 'steps': 'Hire\nExpand\nLaunch EU'}}]}}
    r = db.run(json.dumps(st), skip_film_only=True)
    assert not r['errors'] and r['slides'] == 3 and any(w['code'] == 'film-only' for w in r['warnings'])


def test_device_scene_becomes_bullets_with_warning():
    r = db.run(json.dumps({'scenes': [{'type': 'device', 'f': {'heading': 'Built for *every screen*', 'items': 'Plan\nBuild'}}]}))
    assert not r['errors'] and r['slides'] == 1 and any(w['code'] == 'film-alias' for w in r['warnings'])


CLAUDE_MODS = {'heading': '1. Claude Mods: plugins that change deeper behavior', 'items': [
    "Shipped in 2.1.287 (Oct 1): plugins can now modify Claude Code's deeper behavior, not just add commands",
    'Mods hook events like tool.check, prompt.submit, agent.spawn and turn.step',
    'They can draw their own UI panes, add autocomplete rows and raise native notifications ($.ui.notify)',
    'Org-managed mods take priority, and a user mod that interferes with an org guard is unloaded']}

def test_four_long_items_under_a_two_line_heading_fit_on_both_targets(tmp_path):
    """Seen in a real agent run (2026-10-09): four items of about 100 characters under a two-line heading were refused at 20 pt on the `any` target.
    The layout now falls back to a tighter profile (smaller heading, smaller gaps, down to the 18 pt floor) before it gives up."""
    for target in ('any', 'powerpoint', 'google-slides'):
        r = db.run(json.dumps({'target': target, 'scenes': [{'type': 'bullets', 'f': CLAUDE_MODS, 'notes': 'n'}]}), str(tmp_path / f'{target}.pptx'))
        assert r['ok'] and not r['errors'] and r['lint'] == [], (target, r['errors'], r['lint'])

def test_a_list_that_fits_the_roomy_profile_is_laid_out_as_before(tmp_path):
    r = db.run(json.dumps({'target': 'google-slides', 'scenes': [{'type': 'bullets', 'f': {'heading': 'H', 'items': ['a', 'b', 'c']}, 'notes': 'n'}]}), str(tmp_path / 'x.pptx'))
    assert r['ok']

def test_a_list_too_long_even_for_the_tight_profile_is_still_an_exact_error(tmp_path):
    f = {'heading': 'H', 'items': ['x ' * 100] * 6}
    e = db.run(json.dumps({'target': 'any', 'scenes': [{'type': 'bullets', 'f': f}]}), str(tmp_path / 'x.pptx'))['errors']
    assert 'do not fit on the slide even at 18 pt' in e[0]['message'], e


def test_stat_cards_and_idea_cards_build_clean_and_limits_are_exact(tmp_path):
    deck = {'target': 'google-slides', 'scenes': [
        {'type': 'stats', 'f': {'kicker': 'In numbers', 'stats': [['10', 'releases', '2.1.286 to 2.1.295'], ['1M', 'token context', 'Haiku 5.5'], ['300+', 'bug fixes']]}, 'notes': 'n'},
        {'type': 'cards', 'f': {'heading': 'Three ideas', 'cards': [['Mods', 'Plugins change deeper behavior.'], ['Side agent', 'Flags what you might miss.'], ['Fail-closed hooks', 'A broken guard blocks the action.']]}, 'notes': 'n'}]}
    r = db.run(json.dumps(deck), str(tmp_path / 'c.pptx')); assert r['ok'] and r['lint'] == [], (r['errors'], r['lint'])
    e = db.run(json.dumps({'scenes': [{'type': 'cards', 'f': {'heading': 'H', 'cards': [['t' * 41, 'x']]}}]}))['errors']
    assert e[0]['message'] == 'slide 1 (cards): cards[1] title \'' + 't' * 41 + '\' is 41 characters, the limit is 40', e[0]['message']
    e = db.run(json.dumps({'scenes': [{'type': 'cards', 'f': {'heading': 'H', 'cards': [['t', 'x']] * 5}}]}))['errors']
    assert e[0]['message'] == 'slide 1 (cards): cards has 5 entries, the limit is 4'

def test_stat_sub_line_and_dict_forms():
    r = db.run('<deck><slide layout="stats"><stat value="1" label="a" sub="s"/><notes>n</notes></slide></deck>'); assert not r['errors']
    r = db.run(json.dumps({'scenes': [{'type': 'stats', 'f': {'stats': [{'value': '1', 'label': 'a', 'sub': 's'}]}}]})); assert not r['errors']
    e = db.run(json.dumps({'scenes': [{'type': 'stats', 'f': {'stats': [['1', 'a', 's' * 41]]}}]}))['errors']; assert 'sub is 41 characters, the limit is 40' in e[0]['message']

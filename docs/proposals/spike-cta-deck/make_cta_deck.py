#!/usr/bin/env python3
"""Spike for docs/proposals/presentation-templates.md: one hand-built call-to-action slide for the A2UI Catalog.

Built the way the proposal describes: a real title placeholder, shapes added in reading order, alt text, speaker notes,
brand colours taken from atoms/brand-tokens.yaml (dark theme, oklch converted to sRGB), text sizes at or above an 18 pt floor,
and a text-fit check that uses real font metrics. It then READS THE FILE BACK to lint it and to draw a rough preview PNG,
because this machine has no PowerPoint or LibreOffice. The preview is our own approximation, not how PowerPoint or Slides draw it.

  python make_cta_deck.py <out.pptx> [<preview.png>]        (needs python-pptx, qrcode, pillow)
"""
import sys, math, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.lang import MSO_LANGUAGE_ID
import qrcode
from PIL import Image, ImageDraw, ImageFont

URL = 'https://a2uicatalog.ai'
FONT = 'Roboto'            # brand stack; in Google Slides. Falls back on machines without it (see the fit margin)
FONT_FILES = {False: '/usr/share/fonts/chromeos/roboto/Roboto-Regular.ttf', True: '/usr/share/fonts/chromeos/roboto/Roboto-Bold.ttf'}
FLOOR_PT = 18
# The target is a variable, not a constant. A target profile is data: what the destination can do and how much margin a layout needs there.
#   safety = the share of a box's measured width the layout may use, so a font substituted by the destination (wider) still fits.
#   google-slides: Roboto exists there, so almost no margin (3% covers the import's slightly different text widths, measured at about 0.5%).
#   powerpoint / any: Roboto may be missing on the machine that opens the file, so keep a 12% margin.
TARGETS = {
    'google-slides': dict(safety=0.97, note='fonts: Google Fonts; opened via Drive import'),
    'powerpoint':    dict(safety=0.88, note='fonts: assume Office-safe fallback'),
    'any':           dict(safety=0.88, note='unknown destination: the safe layout'),
}
TARGET = 'any'
SAFETY = TARGETS[TARGET]['safety']
def set_target(name):
    global TARGET, SAFETY
    if name not in TARGETS: raise SystemExit(f'unknown target {name!r}; choose one of {sorted(TARGETS)}')
    TARGET, SAFETY = name, TARGETS[name]['safety']

# ---- brand colours: oklch -> sRGB (atoms/brand-tokens.yaml, dark theme) ----
def oklch(L, C, h):
    a, b = C * math.cos(math.radians(h)), C * math.sin(math.radians(h))
    l_ = (L + 0.3963377774 * a + 0.2158037573 * b) ** 3; m_ = (L - 0.1055613458 * a - 0.0638541728 * b) ** 3; s_ = (L - 0.0894841775 * a - 1.2914855480 * b) ** 3
    rgb = (4.0767416621 * l_ - 3.3077115913 * m_ + 0.2309699292 * s_, -1.2684380046 * l_ + 2.6097574011 * m_ - 0.3413193965 * s_, -0.0041960863 * l_ - 0.7034186147 * m_ + 1.7076147010 * s_)
    g = lambda v: 12.92 * v if v <= 0.0031308 else 1.055 * (max(v, 0) ** (1 / 2.4)) - 0.055
    return tuple(max(0, min(255, round(g(v) * 255))) for v in rgb)
BRAND = dict(bg=oklch(.27, .025, 255), surface=oklch(.33, .025, 255), border=oklch(.42, .02, 255), text=oklch(.95, .01, 255), muted=oklch(.72, .02, 255),
             accent=oklch(.72, .16, 277), accent_ink=oklch(.15, .02, 255), accent2=oklch(.75, .12, 202), qr_bg=(255, 255, 255), qr_ink=oklch(.22, .02, 255))
rgb = lambda c: RGBColor(*c)
def lum(c):
    f = lambda v: v / 12.92 if v <= .03928 else ((v + .055) / 1.055) ** 2.4
    r, g, b = (f(v / 255) for v in c); return .2126 * r + .7152 * g + .0722 * b
def contrast(a, b):
    x, y = sorted((lum(a), lum(b)), reverse=True); return (x + .05) / (y + .05)

# ---- text measurement with the real font ----
def font(pt, bold): return ImageFont.truetype(FONT_FILES[bold], size=int(round(pt * 10)))       # 10 px per pt, so widths are in 1/10 pt
def width_pt(text, pt, bold): return font(pt, bold).getlength(text) / 10
def wrap(text, pt, bold, box_w_pt):
    lines, cur = [], ''
    for w in text.split():
        t = (cur + ' ' + w).strip()
        if width_pt(t, pt, bold) <= box_w_pt or not cur: cur = t
        else: lines.append(cur); cur = w
    return lines + ([cur] if cur else [])

def best_width(text, pt, bold, lo, hi, step=0.05, max_lines=3):
    """Calibration by search: the box width (inches) whose wrap has the fewest lines and the most even last line.
    Greedy wrapping makes a one-word last line (an orphan) common; trying widths and scoring the raggedness removes it.
    The score is taken on the NOMINAL wrap (what the real renderer does); the fallback-font wrap (box x SAFETY, a wider font)
    only has to stay within max_lines, so a substituted font cannot overflow the box."""
    best = None
    w = lo
    while w <= hi + 1e-9:
        lines = wrap(text, pt, bold, w * 72)
        if len(lines) <= max_lines and len(wrap(text, pt, bold, w * 72 * SAFETY)) <= max_lines:
            widths = [width_pt(l, pt, bold) for l in lines]
            last = widths[-1] / max(widths)
            score = (len(lines), -(min(last, 0.9)), -w)         # fewer lines, then a fuller last line, then the narrower box
            if best is None or score < best[0]: best = (score, round(w, 2), lines)
        w += step
    return best[1], best[2]

def fit_text(text, pts, bold, lo, hi, step=0.05, max_lines=3, min_last=0.4):
    """Pick the largest size (from pts) and a box width where the real wrap and the fallback-font wrap (a wider font, box x SAFETY) give the
    SAME number of lines and the last line is not an orphan. Then the layout cannot shift when a font is substituted, and no line is a stub."""
    for pt in pts:
        best = None
        w = lo
        while w <= hi + 1e-9:
            n, f = wrap(text, pt, bold, w * 72), wrap(text, pt, bold, w * 72 * SAFETY)
            if len(n) == len(f) <= max_lines:
                ws = [width_pt(l, pt, bold) for l in n]; last = ws[-1] / max(ws) if len(n) > 1 else 1.0
                if last >= min_last:
                    rag = (max(ws) - min(ws)) / max(ws) if len(n) > 1 else 0.0          # how uneven the lines are
                    score = (round(rag, 2), round(w, 2))
                    if best is None or score < best[0]: best = (score, pt, round(w, 2), n)
            w += step
        if best: return best[1], best[2], best[3]
    return None

def card_geometry(col_top, col_bottom, qr=3.0, pad=0.3, gap=0.12, cap_h=0.34):
    """The card is sized from its contents (padding, QR, gap, caption, padding) and centred on the text column it sits beside."""
    h = pad + qr + gap + cap_h + pad
    top = (col_top + col_bottom) / 2 - h / 2
    return dict(top=top, h=h, qr_top=top + pad, cap_top=top + pad + qr + gap, cap_h=cap_h, pad=pad, qr=qr)

HEAD = 'Give your agent a catalogue of interfaces it can compose'
def build(out):
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[5]                                # "Title Only": the headline is a real title placeholder
    s = prs.slides.add_slide(layout)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BRAND['bg'])
    fits = []

    def text(shape, runs, pt, bold=False, color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, inset=0.0):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(inset); tf.margin_top = tf.margin_bottom = Inches(inset * .6)
        p = tf.paragraphs[0]; p.alignment = align
        for t, c in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(pt); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(c or color)
            r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_US
        full = ''.join(t for t, _ in runs)
        box_w = shape.width / 12700 - 2 * inset * 72; box_h = shape.height / 12700 - 2 * inset * .6 * 72
        lines = wrap(full, pt, bold, box_w * SAFETY); need = len(lines) * pt * 1.2
        fits.append((shape.name, len(lines), round(need), round(box_h), need <= box_h))
        return shape

    def box(shape_type, x, y, w, h, fill=None, line=None, name=''):
        sh = s.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        else: sh.fill.background()
        if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(1.25)
        else: sh.line.fill.background()
        sh.shadow.inherit = False
        return sh
    def tb(x, y, w, h, name): sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name; return sh

    # reading order = z-order = the order shapes are added
    LH = 1.2                                                           # line height as a multiple of the size (calibrated against a real render)
    hp, hw, hlines = fit_text(HEAD, (44, 42, 40, 38, 36), True, 6.8, 8.1)
    SUB = '500+ declarative building blocks. One open schema. Your agent describes the interface; the renderer draws it.'
    sp, sw, slines = fit_text(SUB, (20,), False, 6.4, 8.1)
    pad = 0.08
    hh, sh_ = len(hlines) * hp * LH / 72 + 2 * pad, len(slines) * sp * LH / 72 + 2 * pad
    chip_h, btn_h, g1, g2, g3 = 0.5, 0.85, 0.3, 0.5, 0.3
    col_h = hh + g1 + sh_ + g2 + chip_h + g3 + btn_h
    free_top, free_bot = 1.15, 7.5 - 0.55                              # below the wordmark, and the same margin at the bottom as at the top
    y0 = free_top + ((free_bot - free_top) - col_h) / 2                # the column is centred in the free area
    y_sub, y_chip, y_btn = y0 + hh + g1, y0 + hh + g1 + sh_ + g2, y0 + hh + g1 + sh_ + g2 + chip_h + g3
    text(tb(0.8, 0.55, 5, 0.5, 'Wordmark'), [('A2UI', BRAND['accent']), (' Catalog', BRAND['text'])], 22, True, BRAND['text'])
    t = s.shapes.title; t.name = 'Headline'; t.left, t.top, t.width, t.height = Inches(0.8), Inches(y0), Inches(hw), Inches(hh)
    text(t, [(HEAD, None)], hp, True, BRAND['text'], anchor=MSO_ANCHOR.TOP)
    text(tb(0.8, y_sub, sw, sh_, 'Subhead'), [(SUB, None)], sp, False, BRAND['muted'])
    x = 0.8
    for label in ('Open schema', 'Many surfaces', 'Built for agents'):
        w = width_pt(label, FLOOR_PT, True) / SAFETY / 72 + 0.45      # size the chip from the same margin the fit check uses
        c = box(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_chip, w, chip_h, BRAND['surface'], BRAND['border'], 'Chip: ' + label)
        text(c, [(label, None)], FLOOR_PT, True, BRAND['text'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.08); x += w + 0.2
    btn = box(MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, y_btn, 3.9, btn_h, BRAND['accent'], None, 'Button: Try the catalogue')
    text(btn, [('Try the catalogue', None)], 26, True, BRAND['accent_ink'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.1)
    btn.click_action.hyperlink.address = URL
    text(tb(4.95, y_btn + (btn_h - 0.5) / 2, 3.6, 0.5, 'Address'), [('a2uicatalog.ai', None)], 22, True, BRAND['accent2'], anchor=MSO_ANCHOR.MIDDLE)
    g = card_geometry(y0, y0 + col_h)                                                   # centred on the column from headline top to button bottom
    card = box(MSO_SHAPE.ROUNDED_RECTANGLE, 9.05, g['top'], 3.5, g['h'], BRAND['qr_bg'], None, 'QR card')
    qr = qrcode.QRCode(border=1, box_size=12, error_correction=qrcode.constants.ERROR_CORRECT_M); qr.add_data(URL); qr.make(fit=True)
    img = qr.make_image(fill_color=BRAND['qr_ink'], back_color=BRAND['qr_bg']).convert('RGB'); img.save('/tmp/_cta_qr.png')
    pic = s.shapes.add_picture('/tmp/_cta_qr.png', Inches(9.05 + (3.5 - g['qr']) / 2), Inches(g['qr_top']), Inches(g['qr']), Inches(g['qr'])); pic.name = 'QR code'
    pic._element.nvPicPr.cNvPr.set('descr', 'QR code that opens a2uicatalog.ai')
    text(tb(9.05 + (3.5 - g['qr']) / 2, g['cap_top'], g['qr'], g['cap_h'], 'QR caption'), [('Scan to open', None)], FLOOR_PT, True, BRAND['accent_ink'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    card._element.nvSpPr.cNvPr.set('descr', '')                      # decorative card behind the QR code
    s.notes_slide.notes_text_frame.text = ('Call to action. Say the one thing: the agent describes the interface, the catalogue draws it. '
                                          'Point at the button or the QR code, both open a2uicatalog.ai.')
    cp = prs.core_properties; cp.title = 'A2UI Catalog: call to action (test deck)'; cp.keywords = 'target=' + TARGET; cp.author = 'A2UI Catalog'; cp.language = 'en-US'
    cp.created = cp.modified = __import__('datetime').datetime(2026, 10, 8, 8, 0, 0)         # fixed, so the same input gives the same file
    prs.save(out); normalise_zip(out); return fits

def normalise_zip(path):
    """python-pptx stamps every zip entry with the build time, so two identical decks differ in bytes. Rewrite the archive with fixed
    timestamps and a stable entry order (names are already deterministic) so the same input gives the same file."""
    import zipfile, os
    with zipfile.ZipFile(path) as zin: items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path + '.tmp', 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0)); zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = 0o644 << 16
            zout.writestr(zi, data)
    os.replace(path + '.tmp', path)

def lint(path, fits):
    prs = Presentation(path); W, H = prs.slide_width, prs.slide_height; sl = prs.slides[0]; probs = []
    if not sl.shapes.title or not sl.shapes.title.text_frame.text.strip(): probs.append('no slide title')
    for sh in sl.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W or sh.top + sh.height > H: probs.append(f'{sh.name}: outside the slide')
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size.pt < FLOOR_PT: probs.append(f'{sh.name}: {r.font.size.pt} pt is under the {FLOOR_PT} pt floor')
        if sh.shape_type == 13 and not sh._element.nvPicPr.cNvPr.get('descr'): probs.append(f'{sh.name}: picture without alt text')
    shapes = [sh for sh in sl.shapes]
    rect = lambda sh: (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
    inside = lambda a, b: a[0] >= b[0] and a[1] >= b[1] and a[2] <= b[2] and a[3] <= b[3]
    for i, a in enumerate(shapes):
        for b in shapes[i + 1:]:
            ra, rb = rect(a), rect(b)
            if ra[0] < rb[2] and rb[0] < ra[2] and ra[1] < rb[3] and rb[1] < ra[3] and not (inside(ra, rb) or inside(rb, ra)):
                probs.append(f'{a.name} overlaps {b.name}')
    # balance: margins of the content block, and the card against the column it sits beside
    box_of = lambda names: [sh for sh in shapes if sh.name in names]
    content = [sh for sh in shapes if sh.name != 'Headline' or True]
    L = min(sh.left for sh in content) / 914400; R = (W - max(sh.left + sh.width for sh in content)) / 914400
    T = min(sh.top for sh in content) / 914400; B = (H - max(sh.top + sh.height for sh in content)) / 914400
    if abs(L - R) > 0.15: probs.append(f'left margin {L:.2f} in vs right {R:.2f} in')
    col = [sh for sh in shapes if sh.name in ('Headline', 'Button: Try the catalogue')]; card = [sh for sh in shapes if sh.name == 'QR card'][0]
    off = ((card.top + card.height / 2) - (min(c.top for c in col) + max(c.top + c.height for c in col)) / 2) / 914400
    if abs(off) > 0.06: probs.append(f'QR card is {off:+.2f} in off the centre of its column')
    for sh in shapes:                                                    # an orphan: a last line much shorter than the others
        if sh.has_text_frame and len(sh.text_frame.text) > 40:
            r0 = sh.text_frame.paragraphs[0].runs[0]; ws = [width_pt(l, r0.font.size.pt, bool(r0.font.bold)) for l in wrap(sh.text_frame.text, r0.font.size.pt, bool(r0.font.bold), sh.width / 12700)]
            if len(ws) > 1 and ws[-1] < 0.3 * max(ws): probs.append(f'{sh.name}: last line is {ws[-1]/max(ws):.0%} of the longest (an orphan)')
    margins = dict(left=round(L, 2), right=round(R, 2), top=round(T, 2), bottom=round(B, 2), card_vs_column=round(off, 3))
    probs += [f'{n}: text needs {need} pt of height, box has {have}' for n, _, need, have, ok in fits if not ok]
    pairs = [('headline on background', BRAND['text'], BRAND['bg']), ('subhead on background', BRAND['muted'], BRAND['bg']), ('chip text on chip', BRAND['text'], BRAND['surface']),
             ('button text on button', BRAND['accent_ink'], BRAND['accent']), ('address on background', BRAND['accent2'], BRAND['bg']), ('wordmark accent on background', BRAND['accent'], BRAND['bg']),
             ('QR caption on card', BRAND['accent_ink'], BRAND['qr_bg'])]
    cr = [(n, round(contrast(a, b), 2)) for n, a, b in pairs]; probs += [f'contrast {n}: {v}:1 is under 4.5' for n, v in cr if v < 4.5]
    return probs, cr, margins

def preview(path, out):
    prs = Presentation(path); sl = prs.slides[0]; S = 100 / 914400 * 12700 / 12700   # 1 inch = 100 px
    W, H = int(prs.slide_width / 914400 * 100), int(prs.slide_height / 914400 * 100)
    im = Image.new('RGB', (W, H), BRAND['bg']); d = ImageDraw.Draw(im)
    px = lambda e: e / 914400 * 100
    for sh in sl.shapes:
        x0, y0, x1, y1 = px(sh.left), px(sh.top), px(sh.left + sh.width), px(sh.top + sh.height)
        if sh.shape_type == 13:
            from io import BytesIO; im.paste(Image.open(BytesIO(sh.image.blob)).convert('RGB').resize((int(x1 - x0), int(y1 - y0))), (int(x0), int(y0))); continue
        if sh.shape_type == 1 and sh.fill.type == 1:
            rad = (sh.adjustments[0] if len(sh.adjustments) else 0) * min(x1 - x0, y1 - y0)      # the shape's own corner setting (PowerPoint's default is 1/6 of the short side)
            fc = tuple(sh.fill.fore_color.rgb); d.rounded_rectangle([x0, y0, x1, y1], radius=rad, fill=fc, outline=tuple(sh.line.color.rgb) if sh.line.fill.type == 1 else None, width=1)
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame; ins = px(tf.margin_left); box_w = (x1 - x0) - 2 * ins; y = y0 + px(tf.margin_top)
            runs = [(r.text, r.font.size.pt, bool(r.font.bold), tuple(r.font.color.rgb)) for p in tf.paragraphs for r in p.runs]
            full = ''.join(r[0] for r in runs); pt, bold = runs[0][1], runs[0][2]
            lines = wrap(full, pt, bold, box_w * 72 / 100 * 1.0); lh = pt * 1.2 * 100 / 72
            tot = lh * len(lines)
            if tf.vertical_anchor == MSO_ANCHOR.MIDDLE: y = y0 + ((y1 - y0) - tot) / 2
            fnt = ImageFont.truetype(FONT_FILES[bold], size=int(round(pt * 100 / 72)))
            for ln in lines:
                lw = fnt.getlength(ln); xx = x0 + ins if tf.paragraphs[0].alignment != PP_ALIGN.CENTER else x0 + ((x1 - x0) - lw) / 2
                pos = 0
                for t, _, b, c in runs:                                         # colour per run, left to right
                    seg = ln[pos:pos + len(t)] if len(runs) > 1 else ln
                    if len(runs) > 1 and not seg: break
                    d.text((xx, y), seg, font=fnt, fill=c); xx += fnt.getlength(seg); pos += len(t)
                    if len(runs) == 1: break
                y += lh
    im.save(out)

if __name__ == '__main__':
    if '--target' in sys.argv:
        i = sys.argv.index('--target'); set_target(sys.argv[i + 1]); del sys.argv[i:i + 2]
    print('target:', TARGET, '-', TARGETS[TARGET]['note'], f'(safety {SAFETY})')
    out = sys.argv[1]; fits = build(out); probs, cr, margins = lint(out, fits)
    print('text fit:'); [print(f'  {n:34s} {l} line(s), needs {need} pt of {have} pt  {"ok" if ok else "OVERFLOW"}') for n, l, need, have, ok in fits]
    print('contrast:'); [print(f'  {n:34s} {v}:1') for n, v in cr]
    print('balance (inches):', margins)
    print('lint:', 'clean' if not probs else probs)
    if len(sys.argv) > 2: preview(out, sys.argv[2]); print('preview ->', sys.argv[2])

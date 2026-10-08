#!/usr/bin/env python3
"""Calibrate the local preview against a real render. Give it the PPTX and a screenshot of the slide as PowerPoint or Google Slides drew it;
it aligns the screenshot to the slide, then for every shape in the file compares where ink actually landed in the real render with where the
local preview put it, and prints the offsets in preview pixels (100 px per inch).

  python calibrate_preview.py <deck.pptx> <screenshot.png>          (needs python-pptx, pillow, numpy)

Reading it: a few pixels either way is noise; a consistent sign for one kind of shape is a calibration constant worth applying (for
example a line-height factor); a big offset on one shape is a layout bug the preview hid. Collect several screenshots before trusting a constant."""
import sys, importlib.util
from pathlib import Path
import numpy as np
from PIL import Image
from pptx import Presentation
here = importlib.util.spec_from_file_location('deck', str(Path(__file__).resolve().parents[3] / 'tools/deck/deck_kit.py')); deck = importlib.util.module_from_spec(here); here.loader.exec_module(deck)

def main(pptx, shot_path):
    deck.preview(pptx, '/tmp/_cal_prev.png'); prev = Image.open('/tmp/_cal_prev.png').convert('RGB'); shot = Image.open(shot_path).convert('RGB')
    bg = np.array(prev.getpixel((5, 5)))
    a = np.array(shot).astype(int); same = (np.abs(a - bg).sum(axis=2) < 18)          # the slide area is where the colour is the slide background
    rows = np.where(same.mean(axis=1) > .5)[0]; cols = np.where(same.mean(axis=0) > .5)[0]
    sl = shot.crop((cols.min(), rows.min(), cols.max() + 1, rows.max() + 1)); print(f'slide found in screenshot: {sl.size}, aspect {sl.width / sl.height:.4f} (slide {prev.width / prev.height:.4f})')
    sl = sl.resize(prev.size, Image.LANCZOS); A, B = np.array(prev).astype(int), np.array(sl).astype(int)
    def ink(arr, box):
        x0, y0, x1, y1 = box; m = (np.abs(arr[y0:y1, x0:x1] - bg).sum(axis=2) > 60); ys, xs = np.where(m)
        return None if not len(ys) else (x0 + xs.min(), y0 + ys.min(), x0 + xs.max() + 1, y0 + ys.max() + 1)
    print(f'{"shape":34s} {"dx0":>4s}{"dy0":>4s}{"dx1":>4s}{"dy1":>4s}   kind')
    for sh in Presentation(pptx).slides[0].shapes:
        px = lambda e: int(e / 914400 * 100)
        box = (max(0, px(sh.left) - 14), max(0, px(sh.top) - 14), min(prev.width, px(sh.left + sh.width) + 14), min(prev.height, px(sh.top + sh.height) + 14))
        p, s = ink(A, box), ink(B, box)
        if p and s: print(f'{sh.name[:34]:34s} {s[0]-p[0]:+4d}{s[1]-p[1]:+4d}{s[2]-p[2]:+4d}{s[3]-p[3]:+4d}   {"text" if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_type != 1 else "shape"}')
if __name__ == '__main__': main(sys.argv[1], sys.argv[2])

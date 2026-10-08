#!/usr/bin/env python3
"""deck_build: the mechanical core of the deck pipeline. A typed scene list in, a PPTX and a report of exact errors out.

  python deck_build.py schema                                 print the layouts and their fields as JSON
  python deck_build.py validate <deck.xml|deck.json>          check the input, print every error with its slide and field
  python deck_build.py build <deck.xml|deck.json> -o out.pptx [--target google-slides|powerpoint|any] [--report r.json] [--preview dir] [--skip-film-only]

No model is involved: the same input gives the same bytes. An agent (or a person, or the MCP wrapper) writes the input; this checks it and lays it out.
The scene fields use the same names as the Schemaestro film scenes (eyebrow, headline, sub, kicker, stats "value | label", steps, b1/l1, foot), so one
scene list can drive a film and a deck. `*word*` marks an accent word, as in the films. Film-only scenes (chart, chat, word, morph, device, captions,
image) have no slide layout; with --skip-film-only they are left out and reported as warnings instead of errors.

XML:  <deck title="..." target="google-slides">
        <slide layout="title"><eyebrow>..</eyebrow><headline>..</headline><sub>..</sub><notes>..</notes></slide>
        <slide layout="bullets"><heading>..</heading><item>..</item>...</slide>
        <slide layout="stats"><kicker>..</kicker><stat value="540+" label="atoms"/>...</slide>
        <slide layout="cta"><headline>..</headline><button href="https://..">Label</button><footer>..</footer></slide>
      </deck>
JSON: {"title": "...", "target": "...", "scenes": [{"type": "title", "eyebrow": "...", ...}]}   (film-style: "stats": "v | l\\n..", "steps": "a\\nb")
Needs python-pptx, pillow, defusedxml; the cta layout also needs node and this repo's renderer (the QR code is the schema_qr atom)."""
import sys, json, re, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import deck_kit as k
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.lang import MSO_LANGUAGE_ID

W_IN, H_IN, MX = 13.333, 7.5, 0.8
LH = 1.2
CONTENT_W = W_IN - 2 * MX

# ---- the layouts, as data: field -> kind, whether required, limits. The validator and `schema` both read this table. ----
LAYOUTS = {
    'title':   dict(film='title', fields={'eyebrow': dict(kind='text', max=40), 'headline': dict(kind='text', required=True, max=90),
                                          'sub': dict(kind='text', max=160)}),
    'bullets': dict(film='steps', fields={'heading': dict(kind='text', required=True, max=70),
                                          'items': dict(kind='list', required=True, min=1, max=6, item_max=110)}),
    'stats':   dict(film='stats', fields={'kicker': dict(kind='text', max=40),
                                          'stats': dict(kind='pairs', required=True, min=1, max=4, value_max=10, label_max=40)}),
    'cta':     dict(film='cta', fields={'headline': dict(kind='text', required=True, max=70), 'b1': dict(kind='text', required=True, max=40),
                                        'l1': dict(kind='url'), 'b2': dict(kind='text', max=40), 'l2': dict(kind='url'), 'foot': dict(kind='text', max=60)}),
}
ALIASES = {'steps': 'bullets', 'big-stat': 'stats', 'numbers': 'stats'}
FILM_ONLY = {'chart', 'chat', 'word', 'morph', 'device', 'captions', 'image', 'quote'}
NOTES_MAX = 1200

def schema():
    return {'slide': dict(width_in=W_IN, height_in=H_IN), 'targets': sorted(k.TARGETS), 'aliases': ALIASES, 'film_only': sorted(FILM_ONLY),
            'common': {'notes': dict(kind='text', max=NOTES_MAX, note='speaker notes; recommended on every slide'), 'alt': dict(kind='text', note='alt text for the slide figure, where a layout has one')},
            'layouts': {n: dict(film_scene=d['film'], fields=d['fields']) for n, d in LAYOUTS.items()}}

# ---- parsing: XML or JSON -> {title, target, scenes:[{layout, f, notes, at}]} ----
def err(errs, at, code, msg, **extra): errs.append(dict(slide=at, code=code, message=msg, **extra))
def lines(v): return [l.strip() for l in str(v or '').split('\n') if l.strip()]
def plain(s): return re.sub(r'\*([^*]+)\*', r'\1', s)
def segs(s):
    """'Meet *Studio* now' -> [('Meet ', False), ('Studio', True), (' now', False)] (the films' accent marker)."""
    out, pos = [], 0
    for m in re.finditer(r'\*([^*]+)\*', s):
        if m.start() > pos: out.append((s[pos:m.start()], False))
        out.append((m.group(1), True)); pos = m.end()
    if pos < len(s): out.append((s[pos:], False))
    return out or [('', False)]
def oneline(s): return re.sub(r'\s+', ' ', str(s or '')).strip()

def parse_xml(src, errs):
    from defusedxml import ElementTree as ET
    try: root = ET.fromstring(src)
    except Exception as e: err(errs, None, 'xml', f'not well-formed XML: {e}'); return None
    if root.tag != 'deck': err(errs, None, 'root', f'the root element must be <deck>, found <{root.tag}>'); return None
    deck = dict(title=root.get('title', ''), target=root.get('target'), scenes=[])
    for i, sl in enumerate(root, 1):
        if sl.tag != 'slide': err(errs, None, 'element', f'unexpected <{sl.tag}> inside <deck>; only <slide> is allowed'); continue
        f, notes = {}, ''
        for ch in sl:
            if ch.tag == 'notes': notes = oneline(ch.text)
            elif ch.tag == 'item': f.setdefault('items', []).append(oneline(ch.text))
            elif ch.tag == 'stat': f.setdefault('stats', []).append((oneline(ch.get('value')), oneline(ch.get('label'))))
            elif ch.tag == 'button':
                n = 'b2' if 'b1' in f else 'b1'; f[n] = oneline(ch.text); f['l' + n[1]] = oneline(ch.get('href'))
            elif ch.tag == 'footer': f['foot'] = oneline(ch.text)
            else: f[ch.tag] = oneline(ch.text)
        deck['scenes'].append(dict(layout=sl.get('layout', ''), f=f, notes=notes, at=i))
    return deck

def parse_json(src, errs):
    try: d = json.loads(src)
    except Exception as e: err(errs, None, 'json', f'not valid JSON: {e}'); return None
    sc = d.get('scenes') if isinstance(d, dict) else d
    if not isinstance(sc, list): err(errs, None, 'json', 'expected {"scenes": [...]} or a list of scenes'); return None
    deck = dict(title=(d.get('title', '') if isinstance(d, dict) else ''), target=(d.get('target') if isinstance(d, dict) else None), scenes=[])
    for i, s in enumerate(sc, 1):
        if not isinstance(s, dict): err(errs, i, 'scene', 'a scene must be an object'); continue
        s = dict(s); typ = s.pop('type', s.pop('layout', '')); notes = oneline(s.pop('notes', ''))
        if 'steps' in s and 'items' not in s: s['items'] = s.pop('steps')
        if 'heading' not in s and typ == 'steps' and 'headline' in s: s['heading'] = s.pop('headline')
        if isinstance(s.get('items'), str): s['items'] = lines(s['items'])
        if isinstance(s.get('stats'), str): s['stats'] = [tuple(x.strip() for x in l.split('|', 1)) + ('',) for l in lines(s['stats'])]; s['stats'] = [(a, b) for a, b, *_ in s['stats']]
        deck['scenes'].append(dict(layout=typ, f={kk: (oneline(v) if isinstance(v, str) else v) for kk, v in s.items()}, notes=notes, at=i))
    return deck

def load(path_or_text, errs):
    src = Path(path_or_text).read_text() if len(path_or_text) < 500 and Path(path_or_text).exists() else path_or_text
    return parse_xml(src, errs) if src.lstrip().startswith('<') else parse_json(src, errs)

# ---- validation: every error names the slide, the field and the limit ----
def validate(deck, errs, warns, skip_film_only=False):
    if deck is None: return None
    if deck['target'] is not None and deck['target'] not in k.TARGETS: err(errs, None, 'target', f"unknown target {deck['target']!r}; choose one of {sorted(k.TARGETS)}")
    keep = []
    for sc in deck['scenes']:
        at, lay = sc['at'], ALIASES.get(sc['layout'], sc['layout'])
        if lay not in LAYOUTS:
            if lay in FILM_ONLY and skip_film_only: warns.append(dict(slide=at, code='film-only', message=f'scene {at} ({lay}) is film-only and has no slide layout; left out')); continue
            hint = 'it is a film-only scene; pass --skip-film-only to leave such scenes out' if lay in FILM_ONLY else f'choose one of {sorted(LAYOUTS)}'
            err(errs, at, 'layout', f'slide {at}: unknown layout {sc["layout"]!r}; {hint}'); continue
        sc['layout'] = lay; spec = LAYOUTS[lay]['fields']
        for name in sc['f']:
            if name not in spec and name != 'alt': err(errs, at, 'field', f'slide {at} ({lay}): unknown field {name!r}; fields are {sorted(spec)}', field=name)
        for name, d in spec.items():
            v = sc['f'].get(name)
            if d['kind'] == 'text':
                if not v:
                    if d.get('required'): err(errs, at, 'required', f'slide {at} ({lay}): {name} is required', field=name)
                elif len(plain(v)) > d['max']: err(errs, at, 'too-long', f'slide {at} ({lay}): {name} is {len(plain(v))} characters, the limit is {d["max"]}', field=name)
            elif d['kind'] == 'url':
                if v and not re.match(r'^https?://[^\s]+$', v): err(errs, at, 'url', f'slide {at} ({lay}): {name} must be an http(s) URL, got {v!r}', field=name)
            elif d['kind'] == 'list':
                v = v or []
                if len(v) < d['min']: err(errs, at, 'count', f'slide {at} ({lay}): {name} needs at least {d["min"]} entry, has {len(v)}', field=name)
                elif len(v) > d['max']: err(errs, at, 'count', f'slide {at} ({lay}): {name} has {len(v)} entries, the limit is {d["max"]}', field=name)
                for j, it in enumerate(v, 1):
                    if len(plain(it)) > d['item_max']: err(errs, at, 'too-long', f'slide {at} ({lay}): {name}[{j}] is {len(plain(it))} characters, the limit is {d["item_max"]}', field=name)
            elif d['kind'] == 'pairs':
                v = v or []
                if len(v) < d['min']: err(errs, at, 'count', f'slide {at} ({lay}): {name} needs at least {d["min"]} entry, has {len(v)}', field=name)
                elif len(v) > d['max']: err(errs, at, 'count', f'slide {at} ({lay}): {name} has {len(v)} entries, the limit is {d["max"]}', field=name)
                for j, (a, b) in enumerate(v, 1):
                    if not a: err(errs, at, 'required', f'slide {at} ({lay}): {name}[{j}] has no value', field=name)
                    elif len(a) > d['value_max']: err(errs, at, 'too-long', f'slide {at} ({lay}): {name}[{j}] value {a!r} is {len(a)} characters, the limit is {d["value_max"]}', field=name)
                    if len(b) > d['label_max']: err(errs, at, 'too-long', f'slide {at} ({lay}): {name}[{j}] label is {len(b)} characters, the limit is {d["label_max"]}', field=name)
        if lay == 'cta' and sc['f'].get('b2') and not sc['f'].get('l2'): warns.append(dict(slide=at, code='no-link', message=f'slide {at} (cta): b2 has no l2 link; the button will not open anything'))
        if lay == 'cta' and not sc['f'].get('l1'): warns.append(dict(slide=at, code='no-link', message=f'slide {at} (cta): b1 has no l1 link; the button will not open anything and the slide has no QR code'))
        if len(sc['notes']) > NOTES_MAX: err(errs, at, 'too-long', f'slide {at}: notes are {len(sc["notes"])} characters, the limit is {NOTES_MAX}', field='notes')
        if not sc['notes']: warns.append(dict(slide=at, code='no-notes', message=f'slide {at} ({lay}): no speaker notes'))
        keep.append(sc)
    if not keep and not errs: err(errs, None, 'empty', 'the deck has no slides')
    deck['scenes'] = keep
    return deck

# ---- building ----
class Slide:
    """One slide under construction. Text sizes come from fit_text with the target's safety margin, so a substituted font cannot change the line count."""
    def __init__(self, prs, report):
        self.s = prs.slides.add_slide(prs.slide_layouts[5]); self.report = report; self.fits = []
        self.s.background.fill.solid(); self.s.background.fill.fore_color.rgb = k.rgb(k.BRAND['bg'])
    def text(self, shape, runs, pt, bold=False, color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, inset=0.0):
        tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(inset); tf.margin_top = tf.margin_bottom = Inches(inset * .6)
        p = tf.paragraphs[0]; p.alignment = align
        for t, c in runs:
            r = p.add_run(); r.text = t; r.font.size = Pt(pt); r.font.bold = bold; r.font.name = k.FONT; r.font.color.rgb = k.rgb(c or color)
            r.font.language_id = MSO_LANGUAGE_ID.ENGLISH_US
        full = ''.join(t for t, _ in runs); box_w = shape.width / 12700 - 2 * inset * 72; box_h = shape.height / 12700 - 2 * inset * .6 * 72
        n = len(k.wrap(full, pt, bold, box_w * k.SAFETY)); need = n * pt * LH
        self.fits.append((shape.name, n, round(need), round(box_h), need <= box_h)); return shape
    def tb(self, x, y, w, h, name): sh = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name; return sh
    def box(self, kind, x, y, w, h, fill=None, line=None, name=''):
        sh = self.s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); sh.name = name
        if fill: sh.fill.solid(); sh.fill.fore_color.rgb = k.rgb(fill)
        else: sh.fill.background()
        if line: sh.line.color.rgb = k.rgb(line); sh.line.width = Pt(1.25)
        else: sh.line.fill.background()
        sh.shadow.inherit = False; return sh
    def title(self, name, x, y, w, h): t = self.s.shapes.title; t.name = name; t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(h); return t
    def accent_runs(self, s, base=None, accent=None):
        return [(t, (accent or k.BRAND['accent']) if a else base) for t, a in segs(s)]
    def decorative(self, sh): sh._element.nvSpPr.cNvPr.set('descr', '')
    def alt(self, sh, text): sh._element.nvSpPr.cNvPr.set('descr', text)
    def wordmark(self): self.text(self.tb(MX, 0.55, 5, 0.5, 'Wordmark'), [('A2UI', k.BRAND['accent']), (' Catalog', k.BRAND['text'])], 22, True, k.BRAND['text'])

def heading_fit(text, sizes, hi, max_lines, at, field, errs):
    r = k.fit_text(plain(text), sizes, True, 5.0, hi, max_lines=max_lines)
    if not r: err(errs, at, 'does-not-fit', f'slide {at}: {field} does not fit in {max_lines} lines at {sizes[-1]} pt or larger on this target; shorten it', field=field); return None
    return r

def center_y(heights, top=1.15, bottom=H_IN - 0.55):
    return top + ((bottom - top) - heights) / 2

def lay_title(sl, f, at, errs):
    hr = heading_fit(f['headline'], (66, 60, 54, 48, 44), CONTENT_W, 3, at, 'headline', errs)
    sr = k.fit_text(plain(f['sub']), (26, 24, 22, 20), False, 5.0, 9.5, max_lines=3) if f.get('sub') else None
    if f.get('sub') and not sr: err(errs, at, 'does-not-fit', f'slide {at}: sub does not fit in 3 lines at 20 pt; shorten it', field='sub')
    if not hr or (f.get('sub') and not sr): return
    hp, hw, hl = hr; pad = 0.08; hh = len(hl) * hp * LH / 72 + 2 * pad
    eh = 0.45 if f.get('eyebrow') else 0; sh_ = (len(sr[2]) * sr[0] * LH / 72 + 2 * pad) if sr else 0
    g = 0.3; total = eh + (g if eh else 0) + hh + (g + sh_ if sr else 0); y = center_y(total)
    sl.wordmark()
    if eh:
        e = sl.tb(MX, y, 8, eh, 'Eyebrow'); sl.text(e, [(plain(f['eyebrow']).upper(), None)], k.FLOOR_PT, True, k.BRAND['accent2']); y += eh + g
    t = sl.title('Headline', MX, y, hw, hh); sl.text(t, sl.accent_runs(f['headline'], k.BRAND['text']), hp, True, k.BRAND['text']); y += hh + g
    if sr: sl.text(sl.tb(MX, y, sr[1], sh_, 'Subhead'), [(plain(f['sub']), None)], sr[0], False, k.BRAND['muted'])

def lay_bullets(sl, f, at, errs):
    hr = heading_fit(f['heading'], (44, 40, 36), CONTENT_W, 2, at, 'heading', errs)
    if not hr: return
    hp, hw, hl = hr; pad = 0.08; hh = len(hl) * hp * LH / 72 + 2 * pad; top = 1.15; bottom = H_IN - 0.55
    for pt in (28, 26, 24, 22, 20):                                         # the largest item size at which the whole list fits
        rows = []
        for it in f['items']:
            r = k.fit_text(plain(it), (pt,), False, 4.0, CONTENT_W - 0.6, max_lines=2)
            if not r: rows = None; break
            rows.append(r)
        if rows is None: continue
        gap = 0.28; total = sum(len(r[2]) * pt * LH / 72 + 2 * pad for r in rows) + gap * (len(rows) - 1)
        if hh + 0.45 + total <= bottom - top: break
    else: err(errs, at, 'does-not-fit', f'slide {at}: the {len(f["items"])} items do not fit on the slide at 20 pt; shorten them or split the slide', field='items'); return
    sl.wordmark(); y = top
    t = sl.title('Heading', MX, y, hw, hh); sl.text(t, sl.accent_runs(f['heading'], k.BRAND['text']), hp, True, k.BRAND['text']); y += hh + 0.45
    for i, (it, r) in enumerate(zip(f['items'], rows), 1):
        ih = len(r[2]) * pt * LH / 72 + 2 * pad; dot = 0.16
        d = sl.box(MSO_SHAPE.OVAL, MX, y + (pt * LH / 72 + 2 * pad - dot) / 2, dot, dot, k.BRAND['accent'], None, f'Bullet {i}'); sl.decorative(d)
        sl.text(sl.tb(MX + 0.5, y, r[1], ih, f'Item {i}'), sl.accent_runs(it, k.BRAND['text']), pt, False, k.BRAND['text'])
        y += ih + gap

def lay_stats(sl, f, at, errs):
    st = f['stats']; n = len(st); gap = 0.5; colw = (CONTENT_W - gap * (n - 1)) / n; ok = None
    for pt in (110, 96, 80, 66, 54, 44):
        if all(k.width_pt(v, pt, True) / k.SAFETY <= colw * 72 for v, _ in st): ok = pt; break
    if not ok: err(errs, at, 'does-not-fit', f'slide {at}: a stat value is too wide for a column of {colw:.1f} in even at 44 pt; shorten it or use fewer stats', field='stats'); return
    lab = []
    for j, (v, l) in enumerate(st, 1):
        r = k.fit_text(plain(l), (26, 24, 22, 20), False, 1.5, colw, max_lines=2) if l else None
        if l and not r: err(errs, at, 'does-not-fit', f'slide {at}: stats[{j}] label does not fit in 2 lines in a column of {colw:.1f} in', field='stats'); return
        lab.append(r)
    vh = ok * LH / 72 + 0.16; lh = max((len(r[2]) * r[0] * LH / 72 + 0.16 for r in lab if r), default=0)
    kh = 0.45; total = kh + 0.4 + vh + 0.2 + lh; y = center_y(total); sl.wordmark()
    t = sl.title('Statistics', MX, y, 8, kh); sl.text(t, [((plain(f['kicker']) if f.get('kicker') else 'Key figures').upper(), None)], k.FLOOR_PT, True, k.BRAND['accent2']); y += kh + 0.4
    colors = [k.BRAND['accent'], k.BRAND['accent2'], k.BRAND['text'], k.BRAND['accent']]
    for j, ((v, l), r) in enumerate(zip(st, lab)):
        x = MX + j * (colw + gap)
        sl.text(sl.tb(x, y, colw, vh, f'Value {j + 1}'), [(v, None)], ok, True, colors[j])
        if r: sl.text(sl.tb(x, y + vh + 0.2, colw, lh, f'Label {j + 1}'), [(plain(l), None)], r[0], False, k.BRAND['muted'])

def lay_cta(sl, f, at, errs):
    hr = heading_fit(f['headline'], (66, 60, 54, 48, 44), 7.6 if f.get('l1') else CONTENT_W, 3, at, 'headline', errs)
    if not hr: return
    hp, hw, hl = hr; pad = 0.08; hh = len(hl) * hp * LH / 72 + 2 * pad; bh, g = 0.85, 0.5
    btns = [(f['b1'], f.get('l1'), True)] + ([(f['b2'], f.get('l2'), False)] if f.get('b2') else [])
    sizes = []
    for lab, _, _ in btns:
        w = k.width_pt(plain(lab), 26, True) / k.SAFETY / 72 + 0.8
        if w > 5.2: err(errs, at, 'does-not-fit', f'slide {at}: button {plain(lab)!r} is too wide ({w:.1f} in; the limit is 5.2); shorten it', field='b1'); return
        sizes.append(w)
    col_h = hh + g + bh; y0 = center_y(col_h); sl.wordmark()
    t = sl.title('Headline', MX, y0, hw, hh); sl.text(t, sl.accent_runs(f['headline'], k.BRAND['text']), hp, True, k.BRAND['text'])
    x, yb = MX, y0 + hh + g
    for i, ((lab, href, primary), w) in enumerate(zip(btns, sizes), 1):
        b = sl.box(MSO_SHAPE.ROUNDED_RECTANGLE, x, yb, w, bh, k.BRAND['accent'] if primary else k.BRAND['surface'], None if primary else k.BRAND['border'], f'Button {i}: {plain(lab)}')
        sl.text(b, [(plain(lab), None)], 26, True, k.BRAND['accent_ink'] if primary else k.BRAND['text'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.1)
        if href: b.click_action.hyperlink.address = href
        x += w + 0.3
    if f.get('foot'): sl.text(sl.tb(MX, H_IN - 0.55 - 0.45, 8, 0.45, 'Footer'), [(plain(f['foot']), None)], k.FLOOR_PT, True, k.BRAND['accent2'])
    if f.get('l1'):
        g2 = k.card_geometry(y0, y0 + col_h); cx = W_IN - MX - 3.5
        card = sl.box(MSO_SHAPE.ROUNDED_RECTANGLE, cx, g2['top'], 3.5, g2['h'], k.BRAND['qr_bg'], None, 'QR card'); sl.decorative(card)
        n, cells = k.qr_from_atom(f['l1']); qs = Inches(g2['qr']); rects = k.runs(cells)
        fb = sl.s.shapes.build_freeform(rects[0][0], rects[0][1], scale=qs / n)
        for i, (x0, y1, w) in enumerate(rects):
            if i: fb.move_to(x0, y1)
            fb.add_line_segments([(x0 + w, y1), (x0 + w, y1 + 1), (x0, y1 + 1)], close=True)
        pic = fb.convert_to_shape(Inches(cx + (3.5 - g2['qr']) / 2), Inches(g2['qr_top'])); pic.name = 'QR code (schema_qr atom)'
        pic.fill.solid(); pic.fill.fore_color.rgb = k.rgb(k.BRAND['qr_ink']); pic.line.fill.background(); pic.shadow.inherit = False
        sl.alt(pic, f.get('alt') or f'QR code that opens {f["l1"]}')
        sl.text(sl.tb(cx + (3.5 - g2['qr']) / 2, g2['cap_top'], g2['qr'], g2['cap_h'], 'QR caption'), [('Scan to open', None)], k.FLOOR_PT, True, k.BRAND['accent_ink'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

LAYOUT_FN = {'title': lay_title, 'bullets': lay_bullets, 'stats': lay_stats, 'cta': lay_cta}

def build(deck, out, errs):
    k.set_target(deck['target'] or 'any')
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN); fits = []
    for sc in deck['scenes']:
        sl = Slide(prs, None); n0 = len(errs)
        LAYOUT_FN[sc['layout']](sl, sc['f'], sc['at'], errs)
        if len(errs) > n0: continue
        if sc['notes']: sl.s.notes_slide.notes_text_frame.text = sc['notes']
        fits += [(f'slide {sc["at"]}: {n}', *rest) for n, *rest in sl.fits]
    if errs: return None
    cp = prs.core_properties; cp.title = deck['title'] or 'Untitled deck'; cp.keywords = 'target=' + k.TARGET; cp.author = 'A2UI Catalog'; cp.language = 'en-US'
    cp.created = cp.modified = __import__('datetime').datetime(2026, 10, 8, 8, 0, 0)       # fixed, so the same input gives the same file
    prs.save(out); k.normalise_zip(out); return fits

# ---- lint: read the file back and check what a viewer would hit ----
def lint(path, fits, deck):
    prs = Presentation(path); W, H = prs.slide_width, prs.slide_height; probs = []
    for i, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title
        if i <= len(deck['scenes']) and deck['scenes'][i - 1]['layout'] != 'stats' and (not title or not title.text_frame.text.strip()): probs.append(f'slide {i}: no slide title')
        shapes = list(sl.shapes); rect = lambda s: (s.left, s.top, s.left + s.width, s.top + s.height)
        inside = lambda a, b: a[0] >= b[0] and a[1] >= b[1] and a[2] <= b[2] and a[3] <= b[3]
        for sh in shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W or sh.top + sh.height > H: probs.append(f'slide {i}: {sh.name} is outside the slide')
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt < k.FLOOR_PT: probs.append(f'slide {i}: {sh.name} is {r.font.size.pt} pt, under the {k.FLOOR_PT} pt floor')
            if sh.shape_type == 5 and not sh._element.nvSpPr.cNvPr.get('descr'): probs.append(f'slide {i}: {sh.name} is a vector shape without alt text')
        for a_i, a in enumerate(shapes):
            for b in shapes[a_i + 1:]:
                ra, rb = rect(a), rect(b)
                if ra[0] < rb[2] and rb[0] < ra[2] and ra[1] < rb[3] and rb[1] < ra[3] and not (inside(ra, rb) or inside(rb, ra)): probs.append(f'slide {i}: {a.name} overlaps {b.name}')
        for sh in shapes:                                                    # contrast of every run against the fill of its shape, or the slide background
            if not sh.has_text_frame: continue
            bg = tuple(sh.fill.fore_color.rgb) if sh.shape_type == 1 and sh.fill.type == 1 else None
            for under in shapes[:shapes.index(sh)]:                          # else the fill of the shape it sits on (a card), else the slide background
                if bg is None and under.shape_type == 1 and under.fill.type == 1 and inside(rect(sh), rect(under)): bg = tuple(under.fill.fore_color.rgb)
            bg = bg or k.BRAND['bg']
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.color and r.font.color.type is not None:
                        c = k.contrast(tuple(r.font.color.rgb), bg)
                        if c < 4.5: probs.append(f'slide {i}: {sh.name}: contrast {c:.2f}:1 is under 4.5')
    probs += [f'{n}: text needs {need} pt of height, box has {have}' for n, _, need, have, ok in fits if not ok]
    return sorted(set(probs))

def run(src, out=None, target=None, skip_film_only=False, preview_dir=None):
    """The one entry point (the CLI and the MCP wrapper both call it). Returns a report dict; the PPTX is written to `out` when there are no errors."""
    errs, warns = [], []; deck = load(src, errs); deck = validate(deck, errs, warns, skip_film_only)
    rep = dict(ok=False, errors=errs, warnings=warns)
    if deck is None: return rep
    if target: deck['target'] = target
    if not deck['target']: warns.append(dict(slide=None, code='target', message="no target given; using 'any' (the safe layout with a 12% fallback-font margin). Say google-slides for the tighter layout"))
    rep['target'] = deck['target'] or 'any'; rep['slides'] = len(deck['scenes'])
    if errs or not out: rep['ok'] = not errs; return rep
    fits = build(deck, out, errs)
    if errs: return rep
    rep['lint'] = lint(out, fits, deck); rep['ok'] = not rep['lint']; rep['file'] = str(out)
    import hashlib; rep['sha256'] = hashlib.sha256(Path(out).read_bytes()).hexdigest()
    if preview_dir:
        Path(preview_dir).mkdir(parents=True, exist_ok=True)
        for i in range(len(deck['scenes'])): k.preview(out, str(Path(preview_dir) / f'slide-{i + 1}.png'), i)
        rep['preview'] = str(preview_dir)
    return rep

def main(argv):
    ap = argparse.ArgumentParser(prog='deck_build', description=__doc__.split('\n')[0]); sub = ap.add_subparsers(dest='cmd', required=True)
    sub.add_parser('schema'); v = sub.add_parser('validate'); v.add_argument('input'); v.add_argument('--skip-film-only', action='store_true')
    b = sub.add_parser('build'); b.add_argument('input'); b.add_argument('-o', '--out', required=True); b.add_argument('--target'); b.add_argument('--report'); b.add_argument('--preview'); b.add_argument('--skip-film-only', action='store_true')
    a = ap.parse_args(argv)
    if a.cmd == 'schema': print(json.dumps(schema(), indent=1)); return 0
    rep = run(a.input, getattr(a, 'out', None), getattr(a, 'target', None), a.skip_film_only, getattr(a, 'preview', None))
    if getattr(a, 'report', None): Path(a.report).write_text(json.dumps(rep, indent=1))
    for e in rep['errors']: print('ERROR  ', e['message'])
    for w in rep['warnings']: print('warning', w['message'])
    for p in rep.get('lint', []): print('LINT   ', p)
    print('ok' if rep['ok'] else 'FAILED', f"- {rep.get('slides', 0)} slide(s), target {rep.get('target')}" + (f", {rep['file']}" if 'file' in rep else ''))
    return 0 if rep['ok'] else 1

if __name__ == '__main__': sys.exit(main(sys.argv[1:]))
